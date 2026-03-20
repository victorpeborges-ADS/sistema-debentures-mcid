"""
Converte parecer em Markdown para Word (.docx).
Uso: python md_para_docx.py <arquivo.md> [saida.docx]
Ou: md_para_docx_bytes(md_text) -> bytes para download
"""

import re
import sys
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def add_table_from_md(doc, lines):
    """Cria tabela a partir de linhas markdown | col1 | col2 |."""
    rows = []
    for line in lines:
        line = line.strip()
        if not line or not line.startswith("|"):
            continue
        cells = [c.strip().replace("**", "") for c in line.split("|")[1:-1]]
        if cells and not all(c.replace("-", "").strip() == "" for c in cells):
            rows.append(cells)
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            if j < ncols:
                table.rows[i].cells[j].text = cell


def md_para_docx(md_path: str, docx_path: str = None):
    if docx_path is None:
        docx_path = str(Path(md_path).with_suffix(".docx"))

    with open(md_path, "r", encoding="utf-8") as f:
        content = f.read()

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    lines = content.split("\n")
    i = 0
    table_lines = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=0)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=1)
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=2)
        elif stripped == "---":
            doc.add_paragraph()
        elif stripped.startswith("|"):
            table_lines = [line]
            i += 1
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_table_from_md(doc, table_lines)
            table_lines = []
            continue
        elif stripped.startswith("*") and stripped.endswith("*"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(stripped[1:-1])
            run.italic = True
        elif stripped:
            # Processar negrito **texto**
            p = doc.add_paragraph()
            parts = re.split(r"(\*\*[^*]+\*\*)", stripped)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    run = p.add_run(part[2:-2] + " ")
                    run.bold = True
                else:
                    p.add_run(part)
        i += 1

    doc.save(docx_path)
    return docx_path


def md_para_docx_bytes(md_text: str) -> bytes:
    """Converte texto Markdown para bytes de arquivo .docx."""
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    lines = md_text.split("\n")
    i = 0

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=0)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=1)
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=2)
        elif stripped == "---":
            doc.add_paragraph()
        elif stripped.startswith("|"):
            table_lines = [line]
            i += 1
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_table_from_md(doc, table_lines)
            continue
        elif stripped.startswith("*") and stripped.endswith("*"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(stripped[1:-1])
            run.italic = True
        elif stripped:
            p = doc.add_paragraph()
            parts = re.split(r"(\*\*[^*]+\*\*)", stripped)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    run = p.add_run(part[2:-2] + " ")
                    run.bold = True
                else:
                    p.add_run(part)
        i += 1

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


if __name__ == "__main__":
    md_file = sys.argv[1] if len(sys.argv) > 1 else None
    out_file = sys.argv[2] if len(sys.argv) > 2 else None
    if not md_file:
        print("Uso: python md_para_docx.py <arquivo.md> [saida.docx]")
        sys.exit(1)
    out = md_para_docx(md_file, out_file)
    print(f"Salvo: {out}")
